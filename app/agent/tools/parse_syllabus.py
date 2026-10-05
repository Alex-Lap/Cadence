"""Tool: PDF -> SyllabusExtraction via structured output."""

from __future__ import annotations
from app.schemas import SyllabusExtraction


import io

from pypdf import PdfReader

from app.llm.provider import get_chat_model
from app.schemas import SyllabusExtraction


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"--- page {i} ---\n{text}")
    return "\n\n".join(parts)


def _docx_text(data: bytes) -> str:
    """Paragraphs AND tables.

    python-docx keeps them in separate collections -- document.paragraphs
    silently skips everything inside a table. In a syllabus the assignment
    schedule is almost always a table, so missing them would drop exactly
    the content we care most about.
    """
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for t_i, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"--- table {t_i} ---\n" + "\n".join(rows))

    return "\n\n".join(parts)


def extract_text(data: bytes) -> str:
    """Flatten a syllabus to plain text. Dispatches on magic bytes, not on
    the filename -- users rename things."""
    if data[:4] == b"%PDF":
        return _pdf_text(data)
    if data[:2] == b"PK":          # zip container: .docx, .odt, .pptx
        return _docx_text(data)
    raise ValueError(
        f"Unsupported file type (first bytes: {data[:8]!r}). Expected PDF or DOCX."
    )

SYSTEM_PROMPT = """You extract graded work from a syllabus.

A syllabus usually contains TWO different tables:
  1. A grading policy table -- categories and their weight ("Homework: 30%").
  2. A course schedule -- individual dated deliverables ("HW3, due Oct 14").

Prefer (2). Emit one assignment per individual dated deliverable you find,
and split the category's weight evenly across its members unless the syllabus
says otherwise.

Only emit a row with is_category=true when a category has weight but you can
find no individual dated items for it. Never invent due dates to make a
category look schedulable -- mark it and move on.

Attendance and participation are graded but involve no work to plan: emit
them with is_participation=true and no due date, not as categories.

Never invent an assignment that is not in the document.

DATES vs WEEKS -- this matters:
  * If the syllabus gives an explicit calendar date, put it in due_at.
  * If the syllabus expresses timing in COURSE WEEKS, set due_week to the
    week the item is DUE and leave due_at null. Do NOT convert weeks to
    calendar dates yourself -- code does that downstream.

Also set available_week to the week the item is HANDED OUT. Work cannot begin
before an assignment is assigned, and the scheduler needs that boundary.

Computing due_week: read each item's OWN handout week from the schedule and
add the stated offset. "Project 3 handout in Week 7; due after three weeks"
-> due_week=10. Handout weeks are often irregular (4, 5, 7, 9, 11) -- read
every one from the document. Never assume assignments are evenly spaced, and
never reuse the handout week as the due week.

CLASS MEETINGS: also record when the class regularly meets (lectures, labs,
discussion sections) in `meetings`, from lines like "Lecture: TR 6:00-8:40 PM,
CSE E119". Copy the times exactly as written, converted to 24-hour, even if
they look wrong -- a person checks them afterwards. Put the line you read them
from in source_text. Office hours are NOT meetings. If no meeting times are
given, leave meetings empty; never guess."""


def parse_syllabus(pdf_text: str, *, term_start: str | None = None) -> SyllabusExtraction:
    """Extract courses and graded items from syllabus text."""
    model = get_chat_model(temperature=0)
    structured = model.with_structured_output(SyllabusExtraction)

    context = f"Term start date: {term_start}\n\n" if term_start else ""
    return structured.invoke(
        [
            ("system", SYSTEM_PROMPT),
            ("human", f"{context}Syllabus:\n\n{pdf_text}"),
        ]
    )