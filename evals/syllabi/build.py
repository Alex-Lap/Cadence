"""Build the synthetic test syllabi in this folder.

Every course, person and university here is invented ("Northgate University"),
so these files can live in a public repo. Each one is written to exercise a
specific way extraction goes wrong; the matching answer file in evals/gold/
says what a correct reading is, and why.

Regenerate (needs python-docx and reportlab):

    pip install python-docx reportlab
    python evals/syllabi/build.py
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).resolve().parent

# Boilerplate that real syllabi are mostly made of. It adds the noise a reader
# has to see past; none of it contains graded work or dates.
POLICIES = [
    ("h", "Academic Integrity"),
    ("p", "All work submitted for credit must be your own. Collaboration on concepts is "
          "encouraged, but written answers and code must be produced independently unless an "
          "assignment explicitly allows teams. Suspected violations are referred to the Office "
          "of Student Conduct under the Northgate Honor Code."),
    ("h", "Accessibility"),
    ("p", "Students who need accommodations should register with the Northgate Accessibility "
          "Resource Center and share their letter with the instructor as early as possible in "
          "the term."),
    ("h", "Late Work"),
    ("p", "Unless stated otherwise, late submissions lose 10% per day and are not accepted after "
          "three days. Extensions for documented emergencies are granted on request."),
    ("h", "Course Evaluations"),
    ("p", "You will be invited to complete an anonymous course evaluation near the end of the "
          "term. Feedback is read carefully and used to improve the course."),
]


def synth01():
    """Explicit dates in a schedule table; quizzes are an undated category."""
    return "synth01_algorithms.docx", [
        ("title", "CSC 3410 - Design and Analysis of Algorithms"),
        ("p", "Northgate University - Department of Computer Science - Fall 2026"),
        ("p", "Instructor: Dr. Lena Marsh (fictional)    Email: lmarsh@northgate.example"),
        ("p", "Lectures: Tuesday and Thursday, 2:00 - 3:15 PM, Hollis Hall 204"),
        ("p", "Office hours: Wednesday 10:00 - 11:00 AM, Hollis Hall 318, or by appointment."),
        ("h", "Course Description"),
        ("p", "Asymptotic analysis, divide and conquer, greedy algorithms, dynamic programming, "
              "graph algorithms, network flow, and an introduction to NP-completeness."),
        ("h", "Grading"),
        ("table", [["Component", "Weight"],
                   ["Homework (6 assignments, equally weighted)", "30%"],
                   ["Quizzes (short, in class, unannounced)", "10%"],
                   ["Midterm Exam", "25%"],
                   ["Final Exam", "35%"]]),
        ("h", "Schedule"),
        ("table", [["Week", "Date", "Topic", "Due"],
                   ["1", "Tue 8/25", "Asymptotics", ""],
                   ["2", "Tue 9/1", "Recurrences", "HW 1 due Thu 9/3"],
                   ["3", "Tue 9/8", "Divide and conquer", ""],
                   ["4", "Tue 9/15", "Sorting lower bounds", "HW 2 due Thu 9/17"],
                   ["5", "Tue 9/22", "Greedy algorithms", ""],
                   ["6", "Tue 9/29", "Dynamic programming I", "HW 3 due Thu 10/1"],
                   ["7", "Tue 10/6", "Review", "Midterm Exam Thu 10/8 (in class)"],
                   ["8", "Tue 10/13", "Dynamic programming II", ""],
                   ["9", "Tue 10/20", "Graphs: BFS, DFS", "HW 4 due Thu 10/22"],
                   ["10", "Tue 10/27", "Shortest paths", ""],
                   ["11", "Tue 11/3", "Minimum spanning trees", "HW 5 due Thu 11/5"],
                   ["12", "Tue 11/10", "Network flow", ""],
                   ["13", "Tue 11/17", "NP-completeness", "HW 6 due Thu 11/19"],
                   ["14", "Tue 11/24", "Thanksgiving week - no class Thursday", ""],
                   ["15", "Tue 12/1", "Review", ""]]),
        ("p", "Final Exam: Tuesday, December 8, 2026, 10:00 AM - 12:00 PM, Hollis Hall 204."),
        *POLICIES,
    ]


def synth02():
    """Everything in course weeks with release + offset; project milestones."""
    return "synth02_ml_systems.pdf", [
        ("title", "DS 4200 - Machine Learning Systems"),
        ("p", "Northgate University - School of Data Science - Fall 2026"),
        ("p", "Instructor: Prof. Ravi Castell (fictional). Meetings: Mon/Wed 4:05 - 5:20 PM, "
              "Data Science Building 110."),
        ("h", "How deadlines work"),
        ("p", "Assignments are released on the Monday of the week listed below and are due at the "
              "end of the stated number of weeks after release. Weeks are counted from the first "
              "week of classes, so an assignment released in Week 2 and due two weeks after "
              "release is due at the end of Week 4."),
        ("table", [["Assignment", "Released", "Due"],
                   ["Assignment 1: Data pipelines", "Week 2", "two weeks after release"],
                   ["Assignment 2: Training at scale", "Week 3", "three weeks after release"],
                   ["Assignment 3: Model serving", "Week 6", "two weeks after release"],
                   ["Assignment 4: Monitoring and drift", "Week 9", "three weeks after release"]]),
        ("p", "Course project: teams of two submit a one-page project proposal by the end of "
              "Week 7 and a final report by the end of Week 15. There is a single in-class exam "
              "in Week 11."),
        ("h", "Grading"),
        ("table", [["Component", "Weight"],
                   ["Assignments 1-4 (10% each)", "40%"],
                   ["Project proposal", "5%"],
                   ["Project final report", "25%"],
                   ["Exam", "30%"]]),
        ("h", "Topics by week"),
        ("p", "Week 1 introduction; Week 2 data pipelines; Week 3 distributed training; Week 4 "
              "hyperparameter search; Week 5 evaluation; Week 6 serving; Week 7 batching and "
              "caching; Week 8 feature stores; Week 9 monitoring; Week 10 drift; Week 11 exam; "
              "Weeks 12-15 project work and guest lectures."),
        *POLICIES,
    ]


def synth03():
    """Many dated items in mixed date formats; weekly homework is a category;
    an evening exam time that is not a weekly meeting."""
    return "synth03_physics.pdf", [
        ("title", "PHYS 2101 - General Physics I"),
        ("p", "Northgate University - Fall 2026"),
        ("p", "Lecture: MWF 11:45 AM - 12:35 PM, Science Center 100. "
              "Laboratory: Tuesday 3:00 - 4:50 PM, Science Center B12."),
        ("p", "Instructor: Dr. Tomas Iyer (fictional)."),
        ("h", "Grading"),
        ("table", [["Component", "Weight"],
                   ["Online homework (WebAssign, weekly)", "15%"],
                   ["Lab reports (8, equally weighted)", "20%"],
                   ["Exam 1", "20%"],
                   ["Exam 2", "20%"],
                   ["Final Exam", "25%"]]),
        ("h", "Exams"),
        ("p", "Exam 1 is on Thursday, Sept. 24, from 7:00 to 9:00 PM in Science Center 100. "
              "Exam 2 is on October 29 at the same time and place. The cumulative Final Exam is "
              "on 12/10/2026."),
        ("h", "Laboratory reports"),
        ("table", [["Lab", "Report due"],
                   ["Lab 1 - Measurement and uncertainty", "Sept. 4"],
                   ["Lab 2 - Kinematics", "9/11"],
                   ["Lab 3 - Projectile motion", "September 18th"],
                   ["Lab 4 - Newton's laws", "10/2"],
                   ["Lab 5 - Friction", "Oct 9"],
                   ["Lab 6 - Work and energy", "10/23"],
                   ["Lab 7 - Momentum", "Nov. 6"],
                   ["Lab 8 - Rotation", "11/20"]]),
        ("p", "Reports are due by 11:59 PM on the date listed."),
        *POLICIES,
    ]


def synth04():
    """No dates at all: every item must come back undated."""
    return "synth04_modern_europe.docx", [
        ("title", "HIST 2330 - Modern Europe, 1789 to the Present"),
        ("p", "Northgate University - Department of History - Fall 2026"),
        ("p", "Meeting times and room: see the course portal."),
        ("p", "Instructor: Dr. Clara Benoit (fictional)."),
        ("h", "Assessment"),
        ("table", [["Assessment", "Weight", "Timing"],
                   ["Reading responses (weekly, short)", "20%", "Throughout the term"],
                   ["Midterm Exam", "25%", "Date announced in class"],
                   ["Research Paper", "30%", "Due date TBA"],
                   ["Final Exam", "25%", "Per the university final exam schedule"]]),
        ("h", "Weekly Topics"),
        ("p", "Week 1: The French Revolution. Week 2: Napoleon. Week 3: Restoration and "
              "revolution, 1815-1848. Week 4: Industrialization. Week 5: Nationalism and "
              "unification. Week 6: Empire. Week 7: The road to 1914. Week 8: The Great War. "
              "Week 9: Revolution in Russia. Week 10: The interwar years. Week 11: The Second "
              "World War. Week 12: Cold War Europe. Week 13: 1968. Week 14: 1989 and after."),
        *POLICIES,
    ]


def synth05():
    """Attendance / clicker participation traps, project milestones,
    a lecture plus a lab meeting, and office hours that are not a meeting."""
    return "synth05_cell_biology.docx", [
        ("title", "BIO 3200 - Cell Biology"),
        ("p", "Northgate University - Department of Biology - Fall 2026"),
        ("p", "Lecture: Mon/Wed/Fri 9:35 - 10:25 AM, Life Sciences 150"),
        ("p", "Lab section: Thursday 1:55 - 3:50 PM, Life Sciences 020"),
        ("p", "Office hours: Tuesday 2:00 - 3:00 PM (Dr. Hana Lindqvist, fictional)"),
        ("h", "Grading"),
        ("table", [["Component", "Weight", "Date"],
                   ["Attendance", "5%", "Every lecture"],
                   ["Participation (iClicker questions)", "5%", "Every lecture"],
                   ["Reading quizzes (online, before each lecture)", "10%", "Ongoing"],
                   ["Lab notebook", "10%", "Collected Friday, Dec. 4"],
                   ["Group project: proposal", "5%", "Friday, Sept. 25"],
                   ["Group project: progress report", "10%", "Friday, Oct. 30"],
                   ["Group project: final presentation", "15%", "Wednesday, Dec. 2"],
                   ["Midterm Exam", "20%", "Wednesday, Oct. 14 (in class)"],
                   ["Final Exam", "20%", "Friday, Dec. 11, 9:00 AM"]]),
        ("h", "Course Description"),
        ("p", "Structure and function of the eukaryotic cell: membranes, organelles, the "
              "cytoskeleton, signaling, the cell cycle, and cell death, with an emphasis on the "
              "experimental evidence behind each model."),
        *POLICIES,
    ]


def synth06():
    """Spring term: dates without a year must land in 2027, a break that is
    not an assignment, and a final exam in May."""
    return "synth06_microeconomics.pdf", [
        ("title", "ECON 3010 - Intermediate Microeconomics"),
        ("p", "Northgate University - Spring 2027. Classes begin Monday, January 11."),
        ("p", "Class meets Tuesday and Thursday 9:30 - 10:45 AM in Commerce Hall 12. "
              "Instructor: Prof. Daniel Okoro (fictional)."),
        ("h", "Grading"),
        ("p", "Problem sets (5, equally weighted) 30%. Midterm exam 30%. Final exam 40%."),
        ("h", "Key dates"),
        ("table", [["Item", "Date"],
                   ["Problem Set 1 due", "Friday, Jan 29"],
                   ["Problem Set 2 due", "Friday, Feb 12"],
                   ["Problem Set 3 due", "Friday, Feb 26"],
                   ["Midterm exam (in class)", "Thursday, Mar 4"],
                   ["Spring break - no classes", "Mar 15 - Mar 19"],
                   ["Problem Set 4 due", "Friday, Mar 26"],
                   ["Problem Set 5 due", "Friday, Apr 9"],
                   ["Final exam", "Monday, May 3, 2027, 3:00 - 5:00 PM"]]),
        *POLICIES,
    ]


def synth07():
    """No tables: dates, weights and meeting times only in prose."""
    return "synth07_science_writing.docx", [
        ("title", "ENGL 2100 - Writing About Science"),
        ("p", "Northgate University - Fall 2026"),
        ("p", "We meet Mondays and Wednesdays from 1:00 to 2:15 in the afternoon in Ware Hall "
              "301. I'm Dr. Iris Feld (fictional); my office hours are Thursdays after lunch."),
        ("h", "What you'll write"),
        ("p", "There are three essays. Essay 1, a 1,200-word explainer on a recent discovery, is "
              "due at the start of class in Week 5. Essay 2, a profile of a working scientist, is "
              "due in Week 9. Essay 3, an argument about a contested claim, is due on the Friday "
              "of Week 13. Drafts of each essay are exchanged in a workshop the week before it is "
              "due; workshops count toward participation rather than being graded separately."),
        ("p", "The term ends with a revised Final Portfolio of your best work, due Monday, "
              "December 7."),
        ("h", "How you'll be graded"),
        ("p", "The three essays are worth 20% each. The Final Portfolio is worth 25%. The "
              "remaining 15% is participation, which includes workshops and discussion."),
        *POLICIES,
    ]


SYLLABI = [synth01, synth02, synth03, synth04, synth05, synth06, synth07]


def write_docx(path: Path, blocks) -> None:
    from docx import Document

    doc = Document()
    for kind, content in blocks:
        if kind == "title":
            doc.add_heading(content, level=0)
        elif kind == "h":
            doc.add_heading(content, level=2)
        elif kind == "p":
            doc.add_paragraph(content)
        elif kind == "table":
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            table.style = "Table Grid"
            for r, row in enumerate(content):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    doc.core_properties.author = "Cadence test fixtures"
    doc.save(path)


def write_pdf(path: Path, blocks) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    story = []
    for kind, content in blocks:
        if kind == "title":
            story += [Paragraph(content, styles["Title"]), Spacer(1, 6)]
        elif kind == "h":
            story += [Spacer(1, 6), Paragraph(content, styles["Heading2"])]
        elif kind == "p":
            story += [Paragraph(content, styles["BodyText"]), Spacer(1, 4)]
        elif kind == "table":
            t = Table([[Paragraph(c, styles["BodyText"]) for c in row] for row in content])
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [t, Spacer(1, 6)]
    SimpleDocTemplate(str(path), pagesize=letter, title=path.stem,
                      author="Cadence test fixtures", invariant=1).build(story)


def main() -> None:
    for build in SYLLABI:
        name, blocks = build()
        path = HERE / name
        (write_docx if name.endswith(".docx") else write_pdf)(path, blocks)
        print("wrote", path.name)


if __name__ == "__main__":
    main()
